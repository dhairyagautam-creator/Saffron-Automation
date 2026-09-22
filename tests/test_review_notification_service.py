"""Tests for Review System's automated-email workflow
(app.review_notification_service) -- BM-wise combined-file grouping into
ABM-wise emails, hierarchy-driven recipient resolution, the send
pipeline, and the Phase 2 email-authority machinery (module-wide data
version + shared send-history table) added by the email rework that
generalized the attachment from Coverage-Summary-only to a combined
Opus/Coverage/RGD workbook.

Renamed from test_review_coverage_notification_service.py to match
app/review_notification_service.py's own rename (Q3). Hierarchy lookups
(find_by_employee_code/find_by_employee_name) are monkeypatched to a
fixed in-memory map -- same "stub the lookup, test the grouping"
philosophy as the file this replaces; the recipient/grouping logic itself
is UNCHANGED by this rework (see module docstring), so those tests are
carried over verbatim, only the import target renamed.

build_notification_batch() is exercised with a directly-injected
`bm_files` list (bypassing generate_review_bm_files() and any disk I/O)
-- see that function's own docstring for why this parameter exists.
"""

import app.review_notification_service as notif
from app.review_coverage_service import COVERAGE_REPORT_MONTHS, ComputedBmBlock
from app.review_coverage_service import ROW_LABELS as COVERAGE_ROW_LABELS
from app.review_opus_service import FORMULA_ROWS as OPUS_FORMULA_ROWS
from app.review_opus_service import OPUS_REPORT_MONTHS, ComputedHqBlock
from app.review_opus_service import ROW_LABELS as OPUS_ROW_LABELS
from app.review_rgd_service import IDENTITY_KEYS, SUPPORT_KEYS, VISIT_KEYS
from tests.db_isolation import isolate_database


def _hierarchy_row(code, name, email, abm_code=None, abm_name=None):
    return {
        "employee_code": code, "employee_name": name, "email": email,
        "abm_code": abm_code, "abm_name": abm_name,
    }


def _use_hierarchy(monkeypatch, rows_by_code: dict, rows_by_name: dict | None = None):
    rows_by_name = rows_by_name or {}
    monkeypatch.setattr(notif, "find_by_employee_code", lambda module_key, code: rows_by_code.get(code))
    monkeypatch.setattr(notif, "find_by_employee_name", lambda module_key, name: rows_by_name.get(name, []))


def _bm_file(code, name, path="/tmp/x.xlsx"):
    return {"emp_code": code, "name": name, "file_path": path}


# --- ABM grouping (unchanged recipient workflow) --------------------------

def test_abm_with_four_bms_one_email_four_attachments(monkeypatch):
    hierarchy = {
        "A1": _hierarchy_row("A1", "Amit ABM", "amit@x.com"),
        "B1": _hierarchy_row("B1", "BM One", "b1@x.com", abm_code="A1"),
        "B2": _hierarchy_row("B2", "BM Two", "b2@x.com", abm_code="A1"),
        "B3": _hierarchy_row("B3", "BM Three", "b3@x.com", abm_code="A1"),
        "B4": _hierarchy_row("B4", "BM Four", "b4@x.com", abm_code="A1"),
    }
    _use_hierarchy(monkeypatch, hierarchy)
    bm_files = [_bm_file(c, hierarchy[c]["employee_name"]) for c in ("B1", "B2", "B3", "B4")]

    drafts = notif.build_notification_batch("Xandra", bm_files=bm_files)

    assert len(drafts) == 1
    assert drafts[0]["recipient_email"] == "amit@x.com"
    assert len(drafts[0]["file_paths"]) == 4
    assert len(drafts[0]["bm_names"]) == 4


def test_cross_abm_isolation_never_grouped_together(monkeypatch):
    hierarchy = {
        "A1": _hierarchy_row("A1", "Amit ABM", "amit@x.com"),
        "A2": _hierarchy_row("A2", "Anjali ABM", "anjali@x.com"),
        "B1": _hierarchy_row("B1", "BM One", "b1@x.com", abm_code="A1"),
        "B2": _hierarchy_row("B2", "BM Two", "b2@x.com", abm_code="A2"),
    }
    _use_hierarchy(monkeypatch, hierarchy)
    bm_files = [_bm_file("B1", "BM One"), _bm_file("B2", "BM Two")]

    drafts = notif.build_notification_batch("Xandra", bm_files=bm_files)

    assert len(drafts) == 2
    by_recipient = {d["recipient_email"]: d for d in drafts}
    assert by_recipient["amit@x.com"]["bm_names"] == ["BM One"]
    assert by_recipient["anjali@x.com"]["bm_names"] == ["BM Two"]


def test_recipient_correctness_matches_hierarchy(monkeypatch):
    hierarchy = {
        "A1": _hierarchy_row("A1", "Amit ABM", "amit@x.com"),
        "B1": _hierarchy_row("B1", "BM One", "b1@x.com", abm_code="A1"),
    }
    _use_hierarchy(monkeypatch, hierarchy)
    drafts = notif.build_notification_batch("Xandra", bm_files=[_bm_file("B1", "BM One")])
    assert drafts[0]["recipient_name"] == "Amit ABM"
    assert drafts[0]["recipient_email"] == "amit@x.com"
    assert drafts[0]["subject"] == "Saffron Automation - Review Summary (Xandra, 1 BM)"


def test_no_attachment_duplicated_across_drafts(monkeypatch):
    hierarchy = {
        "A1": _hierarchy_row("A1", "Amit ABM", "amit@x.com"),
        "B1": _hierarchy_row("B1", "BM One", "b1@x.com", abm_code="A1"),
    }
    _use_hierarchy(monkeypatch, hierarchy)
    drafts = notif.build_notification_batch("Xandra", bm_files=[_bm_file("B1", "BM One", path="/tmp/b1.xlsx")])
    all_paths = [p for d in drafts for p in d["file_paths"]]
    assert all_paths == ["/tmp/b1.xlsx"]


def test_bm_missing_from_hierarchy_is_unresolved_not_dropped_silently(monkeypatch):
    _use_hierarchy(monkeypatch, {})
    drafts = notif.build_notification_batch("Xandra", bm_files=[_bm_file("B1", "Ghost BM")])
    assert drafts == []


def test_vacant_or_missing_abm_mapping_is_unresolved(monkeypatch):
    hierarchy = {"B1": _hierarchy_row("B1", "BM One", "b1@x.com", abm_code=None, abm_name=None)}
    _use_hierarchy(monkeypatch, hierarchy)
    drafts = notif.build_notification_batch("Xandra", bm_files=[_bm_file("B1", "BM One")])
    assert drafts == []


def test_abm_with_no_email_is_unresolved(monkeypatch):
    hierarchy = {
        "A1": _hierarchy_row("A1", "Amit ABM", ""),
        "B1": _hierarchy_row("B1", "BM One", "b1@x.com", abm_code="A1"),
    }
    _use_hierarchy(monkeypatch, hierarchy)
    drafts = notif.build_notification_batch("Xandra", bm_files=[_bm_file("B1", "BM One")])
    assert drafts == []


def test_duplicate_hierarchy_records_uses_first_match(monkeypatch):
    first = _hierarchy_row("A1", "Amit ABM", "first@x.com")
    second = _hierarchy_row("A2", "Amit ABM", "second@x.com")
    hierarchy_by_code = {"B1": _hierarchy_row("B1", "BM One", "b1@x.com", abm_code=None, abm_name="Amit ABM")}
    _use_hierarchy(monkeypatch, hierarchy_by_code, rows_by_name={"Amit ABM": [first, second]})
    drafts = notif.build_notification_batch("Xandra", bm_files=[_bm_file("B1", "BM One")])
    assert drafts[0]["recipient_email"] == "first@x.com"


def test_empty_bm_files_produces_no_drafts_no_error(monkeypatch):
    _use_hierarchy(monkeypatch, {})
    assert notif.build_notification_batch("Xandra", bm_files=[]) == []


# --- build_notification_batch_all_divisions (Q2) ---------------------------

def test_all_divisions_batch_concatenates_every_division(monkeypatch):
    """Send Emails now sends for every division in one action -- verify
    the batch builder calls build_notification_batch once per division
    (in DIVISIONS order) and concatenates the results, never dropping or
    duplicating a division's own drafts."""
    calls = []

    def fake_build(division, bm_files=None):
        calls.append(division)
        return [{"division": division, "recipient_email": f"{division}@x.com"}]

    monkeypatch.setattr(notif, "build_notification_batch", fake_build)

    drafts = notif.build_notification_batch_all_divisions()

    assert calls == list(notif.DIVISIONS)
    assert [d["division"] for d in drafts] == list(notif.DIVISIONS)


# --- Send pipeline (Q9: no more local history-table writes) ---------------

def test_send_notification_batch_attaches_correct_files_and_records_shared_history(monkeypatch, tmp_path):
    sent = []
    recorded = []

    def fake_open_connection(sender_email, app_password):
        return object(), sender_email

    def fake_send_via_connection(connection, sender_email, to_address, subject, html_body, text_body=None, attachments=None):
        sent.append({"to": to_address, "attachments": attachments})

    monkeypatch.setattr(notif, "get_settings", lambda: {"sender_email": "s@x.com", "app_password": "pw"})
    monkeypatch.setattr(notif, "open_smtp_connection", fake_open_connection)
    monkeypatch.setattr(notif, "send_via_connection", fake_send_via_connection)
    monkeypatch.setattr(notif, "get_data_version", lambda module: 7)
    monkeypatch.setattr(notif, "record_send", lambda module, data_version: recorded.append((module, data_version)))

    file1 = tmp_path / "Review Summary - BM One.xlsx"
    file1.write_bytes(b"fake-xlsx-bytes")

    drafts = [{
        "recipient_name": "Amit ABM", "recipient_email": "amit@x.com", "division": "Xandra",
        "bm_names": ["BM One"], "file_paths": [str(file1)],
        "subject": "Saffron Automation - Review Summary (Xandra, 1 BM)",
        "body": "<html></html>", "text_body": "text", "status": notif.STATUS_DRAFT,
    }]

    result = notif.send_notification_batch(drafts)

    assert result["sent_count"] == 1
    assert result["failed_count"] == 0
    assert len(sent) == 1
    assert sent[0]["to"] == "amit@x.com"
    [(filename, data)] = sent[0]["attachments"]
    assert filename == "Review Summary - BM One.xlsx"
    assert data == b"fake-xlsx-bytes"

    # Q9: the local review_coverage_email_notifications table is no longer
    # written to -- the shared email_send_history table is the sole
    # record, via record_send("review_system", <current data_version>).
    assert recorded == [("review_system", "7")]


def test_send_notification_batch_records_failure_without_stopping_batch_and_no_history_record(monkeypatch, tmp_path):
    recorded = []
    monkeypatch.setattr(notif, "get_settings", lambda: {"sender_email": "s@x.com", "app_password": "pw"})
    monkeypatch.setattr(notif, "open_smtp_connection", lambda sender_email, app_password: (object(), sender_email))
    monkeypatch.setattr(notif, "record_send", lambda module, data_version: recorded.append((module, data_version)))

    def failing_send(*args, **kwargs):
        raise RuntimeError("smtp boom")

    monkeypatch.setattr(notif, "send_via_connection", failing_send)

    file1 = tmp_path / "Review Summary - BM One.xlsx"
    file1.write_bytes(b"data")
    drafts = [{
        "recipient_name": "Amit ABM", "recipient_email": "amit@x.com", "division": "Xandra",
        "bm_names": ["BM One"], "file_paths": [str(file1)],
        "subject": "s", "body": "b", "text_body": "t", "status": notif.STATUS_DRAFT,
    }]

    result = notif.send_notification_batch(drafts)
    assert result["sent_count"] == 0
    assert result["failed_count"] == 1
    assert drafts[0]["status"] == notif.STATUS_FAILED
    assert "smtp boom" in drafts[0]["error_message"]
    # A run with zero successful sends never records a "send happened"
    # entry in the shared history -- matching Path Validator/Work
    # Distribution's own convention.
    assert recorded == []


def test_get_recent_notifications_still_readable_but_never_written_by_send(monkeypatch):
    """Q9: get_recent_notifications() is kept for reading whatever rows
    were recorded before this rework -- confirm it still works against an
    isolated DB with zero rows (never populated by send_notification_batch
    anymore)."""
    isolate_database(monkeypatch)
    assert notif.get_recent_notifications() == []


# --- Send Emails button state / module-wide data version (Q7) -------------

def test_send_button_state_no_data():
    assert notif.send_button_state(has_data=False, changed_since_last_send=True) == "no_data"


def test_send_button_state_new_data():
    assert notif.send_button_state(has_data=True, changed_since_last_send=True) == "new_data"


def test_send_button_state_resend_confirm():
    assert notif.send_button_state(has_data=True, changed_since_last_send=False) == "resend_confirm"


def test_all_divisions_data_state_no_data_ever_generated(monkeypatch):
    monkeypatch.setattr(notif, "get_data_version", lambda module: 0)
    has_data, changed, last_send = notif.all_divisions_data_state()
    assert (has_data, changed, last_send) == (False, False, None)


def test_all_divisions_data_state_changed_since_last_send(monkeypatch):
    monkeypatch.setattr(notif, "get_data_version", lambda module: 5)
    monkeypatch.setattr(notif, "get_last_send", lambda module: {"data_version": "4", "sent_at": None, "sent_by_name": "X"})
    has_data, changed, last_send = notif.all_divisions_data_state()
    assert has_data is True
    assert changed is True


def test_all_divisions_data_state_not_changed_since_last_send(monkeypatch):
    monkeypatch.setattr(notif, "get_data_version", lambda module: 5)
    monkeypatch.setattr(notif, "get_last_send", lambda module: {"data_version": "5", "sent_at": None, "sent_by_name": "X"})
    has_data, changed, last_send = notif.all_divisions_data_state()
    assert has_data is True
    assert changed is False


# --- all_divisions_ready (Q5) ----------------------------------------------

def test_all_divisions_ready_true_only_when_every_division_ready(monkeypatch):
    monkeypatch.setattr(notif, "_division_reports_ready", lambda division: True)
    assert notif.all_divisions_ready() is True


def test_all_divisions_ready_false_if_any_single_division_not_ready(monkeypatch):
    ready_map = {d: True for d in notif.DIVISIONS}
    ready_map[notif.DIVISIONS[-1]] = False
    monkeypatch.setattr(notif, "_division_reports_ready", lambda division: ready_map[division])
    assert notif.all_divisions_ready() is False


def test_division_reports_ready_false_if_opus_mapping_missing(monkeypatch):
    monkeypatch.setattr(notif, "coverage_prerequisites_ready", lambda division: (True, []))
    monkeypatch.setattr(notif, "OPUS_HQ_BLOCKS_BY_DIVISION", {})  # no mapping for any division
    monkeypatch.setattr(notif, "opus_prerequisites_ready", lambda division: (True, []))
    monkeypatch.setattr(notif, "rgd_prerequisites_ready", lambda division: (True, []))
    assert notif._division_reports_ready("Xandra") is False


def test_division_reports_ready_false_if_rgd_not_ready(monkeypatch):
    monkeypatch.setattr(notif, "coverage_prerequisites_ready", lambda division: (True, []))
    monkeypatch.setattr(notif, "OPUS_HQ_BLOCKS_BY_DIVISION", {"Xandra": ("block",)})
    monkeypatch.setattr(notif, "opus_prerequisites_ready", lambda division: (True, []))
    monkeypatch.setattr(notif, "rgd_prerequisites_ready", lambda division: (False, ["coverage_visits_support_xandra"]))
    assert notif._division_reports_ready("Xandra") is False


# --- HQ resolution (Coverage's bare "Reporting HQ" -> Opus's own spelling) --

def test_resolve_opus_block_exact_match():
    block = object()
    by_hq = {"GUNTUR": block}
    assert notif._resolve_opus_block_for_bm(by_hq, "Guntur") is block


def test_resolve_opus_block_pool_suffix_match():
    block = object()
    by_hq = {"AHMEDABAD POOL": block}
    assert notif._resolve_opus_block_for_bm(by_hq, "Ahmedabad") is block


def test_resolve_opus_block_via_spelling_alias():
    block = object()
    by_hq = {"CHHAPRA": block}  # Annual Targets' own canonical spelling
    assert notif._resolve_opus_block_for_bm(by_hq, "Chapra") is block  # Avg & Calls' bare spelling


def test_resolve_opus_block_returns_none_when_unmatched():
    assert notif._resolve_opus_block_for_bm({"SOMEWHERE ELSE": object()}, "Nowhere") is None


# --- Combined per-BM file generation (3 sheets, order Opus/Coverage/RGD) --

def _coverage_block(emp_code, name, hq="Guntur"):
    rows = {label: {m: 1.0 for m in COVERAGE_REPORT_MONTHS} for label in COVERAGE_ROW_LABELS}
    return ComputedBmBlock(
        division="Xandra", region="Andhra Pradesh", hq=hq, emp_code=emp_code,
        name=name, designation="BM", rows=rows,
    )


def _opus_block(hq="GUNTUR"):
    non_formula_labels = [label for label in OPUS_ROW_LABELS if label not in OPUS_FORMULA_ROWS]
    source_rows = {label: {m: 1.0 for m in OPUS_REPORT_MONTHS} for label in non_formula_labels}
    return ComputedHqBlock(region="Andhra Pradesh", hq=hq, unresolved=False, no_of_bm=1, source_rows=source_rows)


def _rgd_row(bm_code, bm_name, dr_code):
    row = {key: "" for key in IDENTITY_KEYS + SUPPORT_KEYS + VISIT_KEYS}
    row.update({"bm_code": bm_code, "bm_name": bm_name, "dr_code": dr_code, "region": "Andhra Pradesh", "hq": "Guntur"})
    return row


def _apply_generation_fixture(monkeypatch, coverage_blocks, opus_by_hq, rgd_rows, tmp_path):
    monkeypatch.setattr(notif, "_division_reports_ready", lambda division: True)
    monkeypatch.setattr(notif, "_compute_coverage_blocks", lambda division: coverage_blocks)
    monkeypatch.setattr(notif, "compute_opus_blocks_by_hq", lambda division: opus_by_hq)
    monkeypatch.setattr(notif, "_load_rgd_rows", lambda division: rgd_rows)
    monkeypatch.setattr(notif, "REVIEW_UPLOADS_DIR", tmp_path)


def test_combined_file_has_three_sheets_in_opus_coverage_rgd_order(monkeypatch, tmp_path):
    import openpyxl

    coverage_blocks = [_coverage_block("E1", "Alice BM")]
    opus_by_hq = {"GUNTUR": _opus_block()}
    rgd_rows = [_rgd_row("E1", "Alice BM", "D1"), _rgd_row("E2", "Other BM", "D2")]
    _apply_generation_fixture(monkeypatch, coverage_blocks, opus_by_hq, rgd_rows, tmp_path)

    result = notif.generate_review_bm_files("Xandra")
    assert result["success"] is True, result["errors"]
    [f] = result["files"]

    wb = openpyxl.load_workbook(f["file_path"])
    assert wb.sheetnames == ["OPUS SUMMARY", "COVERAGE SUMMARY", "RGD VISIT AND SUPPORT"]

    coverage_rows = list(wb["COVERAGE SUMMARY"].iter_rows(min_row=2, values_only=True))
    assert len(coverage_rows) == len(COVERAGE_ROW_LABELS)  # exactly this one BM's own 9-row block

    rgd_rows_written = list(wb["RGD VISIT AND SUPPORT"].iter_rows(min_row=2, values_only=True))
    assert len(rgd_rows_written) == 1  # only Alice's own row, Other BM's excluded


def test_combined_file_unresolved_opus_hq_writes_placeholder_not_crash(monkeypatch, tmp_path):
    """A BM whose HQ has no matching Opus Summary section still gets a
    file -- its Opus sheet carries an explanatory placeholder instead of
    a data block (Q5 gap this rework explicitly documents)."""
    import openpyxl

    coverage_blocks = [_coverage_block("E1", "Alice BM", hq="Nowhere HQ")]
    opus_by_hq = {"SOMEWHERE ELSE": _opus_block()}  # never matches "Nowhere HQ"
    _apply_generation_fixture(monkeypatch, coverage_blocks, opus_by_hq, [], tmp_path)

    result = notif.generate_review_bm_files("Xandra")
    assert result["success"] is True, result["errors"]
    [f] = result["files"]

    wb = openpyxl.load_workbook(f["file_path"])
    opus_ws = wb["OPUS SUMMARY"]
    assert opus_ws.cell(row=1, column=1).value is not None
    assert "No Opus Summary section" in opus_ws.cell(row=1, column=1).value


def test_generation_blocked_when_not_all_divisions_reports_ready(monkeypatch, tmp_path):
    monkeypatch.setattr(notif, "_division_reports_ready", lambda division: False)
    monkeypatch.setattr(notif, "REVIEW_UPLOADS_DIR", tmp_path)

    result = notif.generate_review_bm_files("Xandra")
    assert result["success"] is False
    assert result["files"] == []
    assert result["errors"]


def test_generation_unknown_division_reports_error_not_exception(tmp_path, monkeypatch):
    monkeypatch.setattr(notif, "REVIEW_UPLOADS_DIR", tmp_path)
    result = notif.generate_review_bm_files("Nonexistent")
    assert result["success"] is False
    assert result["files"] == []
    assert result["errors"]


def test_full_replace_removes_stale_files_from_previous_run(monkeypatch, tmp_path):
    opus_by_hq = {"GUNTUR": _opus_block()}
    _apply_generation_fixture(
        monkeypatch,
        [_coverage_block("E1", "Alice BM"), _coverage_block("E2", "Bob BM")],
        opus_by_hq, [], tmp_path,
    )
    first = notif.generate_review_bm_files("Xandra")
    assert len(first["files"]) == 2

    _apply_generation_fixture(monkeypatch, [_coverage_block("E1", "Alice BM")], opus_by_hq, [], tmp_path)
    second = notif.generate_review_bm_files("Xandra")
    assert len(second["files"]) == 1

    remaining = list(notif._bm_files_output_dir("Xandra").glob("*.xlsx"))
    assert len(remaining) == 1
